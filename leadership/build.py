import re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
doc=Document()
st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(12); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); st.paragraph_format.line_spacing=1.5
for n in ['Heading 1','Heading 2','Heading 3','Title']:
    s=doc.styles[n]; s.font.name='Arial'; s.font.size=Pt(14 if n in('Heading 1','Title') else 12); s.font.bold=True; s.font.color.rgb=None; s.paragraph_format.line_spacing=1.5
def runs(p,text):
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text):
        if not part: continue
        if part.startswith('**'): r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('*'): r=p.add_run(part[1:-1]); r.italic=True
        else: r=p.add_run(part)
        r.font.name='Arial'; r.font.size=Pt(12)
def add_md(path):
    lines=open(path).read().split('\n'); i=0
    while i<len(lines):
        l=lines[i]
        if l.startswith('# '): doc.add_heading(l[2:],1)
        elif l.startswith('### '): doc.add_heading(l[4:],3)
        elif l.startswith('## '): doc.add_heading(l[3:],2)
        elif re.match(r'\d+\. ',l): runs(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',l))
        elif l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.match(r'\|[-| ]+\|$',lines[i]): rows.append([c.strip() for c in lines[i].strip('|').split('|')])
                i+=1
            t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'
            for r,row in enumerate(rows):
                for c,cell in enumerate(row): runs(t.cell(r,c).paragraphs[0],('**%s**'%cell) if r==0 else cell)
            continue
        elif l.strip(): runs(doc.add_paragraph(),l)
        i+=1
doc.add_paragraph('PART A: Reflection and Plan (submission text)',style='Title')
runs(doc.add_paragraph(),'SRN: [YOU SUPPLY: Student Reference Number only; no name]')
runs(doc.add_paragraph(),'Module: Leadership in Practice, BSc Applied Management')
runs(doc.add_paragraph(),'Note: submit Part A only, after rewriting in your own words and removing all [YOU SUPPLY] notes. Delete Part B. See the AI-use rule in Part B, section 0.')
for f in ['task1.md','task2.md','refs.md']: add_md(f)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
add_md('report.md')
doc.save('Leadership Reflection and Plan - Complete.docx')
