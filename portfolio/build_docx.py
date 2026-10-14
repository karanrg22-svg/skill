import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

doc=Document()
st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
st.paragraph_format.line_spacing=1.5
for n in ['Heading 1','Heading 2','Title']:
    s=doc.styles[n]; s.font.name='Arial'; s.font.size=Pt(12 if n!='Heading 1' else 14); s.font.bold=True
    s.font.color.rgb=None; s.element.rPr.rFonts.set(qn('w:asciiTheme'),'') if False else None
    s.paragraph_format.line_spacing=1.5

def runs(p,text):
    for i,part in enumerate(re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',text)):
        if not part: continue
        if part.startswith('**'): r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('*'): r=p.add_run(part[1:-1]); r.italic=True
        else: r=p.add_run(part)
        r.font.name='Arial'; r.font.size=Pt(12)

doc.add_paragraph('Professional Portfolio',style='Title')
for l in ['Name: Apsara Sunuwar','Role: [YOU SUPPLY: programme / role]','Date: [YOU SUPPLY: submission date]','Title: Building the conditions for independent judgement: identity, an AI-supported placement case, and a career plan']:
    runs(doc.add_paragraph(),l)

def add_md(path):
    lines=open(path).read().split('\n'); i=0
    while i<len(lines):
        l=lines[i]
        if l.startswith('# '): doc.add_heading(l[2:],1)
        elif l.startswith('## '): doc.add_heading(l[3:],2)
        elif l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.match(r'\|[-| ]+\|$',lines[i]): rows.append([c.strip() for c in lines[i].strip('|').split('|')])
                i+=1
            t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'
            for r,row in enumerate(rows):
                for c,cell in enumerate(row):
                    p=t.cell(r,c).paragraphs[0]; runs(p,('**%s**'%cell) if r==0 and not cell.startswith('**') else cell)
            continue
        elif l.strip(): runs(doc.add_paragraph(),l)
        i+=1
for f in ['section1.md','section2.md','section3.md']: add_md(f)
doc.add_heading('References',1)
for l in open('references.md').read().split('\n')[1:]:
    if l.strip():
        p=doc.add_paragraph(); runs(p,l)
doc.save('Professional Portfolio.docx')
